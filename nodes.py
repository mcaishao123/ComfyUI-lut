"""
ComfyUI LRTemplate / XMP / CUBE Preset Nodes

Provides nodes for loading and applying image presets/filters.
"""

import os
import io
import time
import traceback

import numpy as np
import torch
from PIL import Image as PILImage

from .lr_parser import parse_lrtemplate, get_settings, get_preset_name, list_presets
from .lr_adjust import apply_preset
from .xmp_parser import parse_xmp, list_xmp_presets
from .cube_parser import parse_cube, apply_cube_lut, list_cube_files


# Presets directory
PRESETS_DIR = os.path.join(os.path.dirname(__file__), "presets")

# Supported preset extensions (for Lightroom-style adjustments)
_LR_EXTENSIONS = (".lrtemplate", ".xmp")


# =============================================================================
# Helper: load any LR-style preset file (.lrtemplate or .xmp)
# =============================================================================

def _load_lr_preset(filepath: str) -> dict:
    """Load a .lrtemplate or .xmp file and return a unified preset dict."""
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".lrtemplate":
        parsed = parse_lrtemplate(filepath)
        settings = get_settings(parsed)
        name = get_preset_name(parsed)
    elif ext == ".xmp":
        parsed = parse_xmp(filepath)
        settings = parsed.get("settings", {})
        name = parsed.get("name", os.path.splitext(os.path.basename(filepath))[0])
    else:
        raise ValueError(f"Unsupported preset format: {ext}")

    return {"name": name, "settings": settings, "source": filepath}


def _list_lr_presets(directory: str) -> list:
    """List all .lrtemplate and .xmp files in a directory."""
    result = []
    if not os.path.isdir(directory):
        return result
    for fname in sorted(os.listdir(directory)):
        if any(fname.lower().endswith(ext) for ext in _LR_EXTENSIONS):
            result.append((fname, os.path.join(directory, fname)))
    return result


# =============================================================================
# LRTemplate / XMP Preset Nodes
# =============================================================================

class LRTemplateLoader:
    """
    Load and parse a .lrtemplate or .xmp preset file.
    Place preset files in the 'presets' subfolder of this node.
    """

    @classmethod
    def INPUT_TYPES(cls):
        preset_list = cls._get_preset_list()
        if not preset_list:
            preset_list = ["No presets found - add .lrtemplate or .xmp files to presets folder"]

        return {
            "required": {
                "preset": (preset_list, {"default": preset_list[0]}),
            },
        }

    RETURN_TYPES = ("LR_PRESET",)
    RETURN_NAMES = ("preset",)
    FUNCTION = "load_preset"
    CATEGORY = "image/filter"

    @classmethod
    def _get_preset_list(cls):
        """Scan the presets directory for .lrtemplate and .xmp files."""
        presets = _list_lr_presets(PRESETS_DIR)
        return [name for name, _ in presets] if presets else []

    @classmethod
    def IS_CHANGED(cls, preset):
        filepath = os.path.join(PRESETS_DIR, preset)
        if os.path.exists(filepath):
            return os.path.getmtime(filepath)
        return float("nan")

    def load_preset(self, preset):
        filepath = os.path.join(PRESETS_DIR, preset)

        if not os.path.exists(filepath):
            print(f"[LRTemplate] Preset file not found: {filepath}")
            return ({},)

        try:
            result = _load_lr_preset(filepath)
            print(f"[LRTemplate] Loaded preset: {result['name']} ({len(result['settings'])} settings)")
            return (result,)
        except Exception as e:
            print(f"[LRTemplate] Error parsing preset '{preset}': {e}")
            return ({},)


class LRTemplateApply:
    """
    Apply a Lightroom preset (.lrtemplate / .xmp) to an image.
    Supports blending with a strength parameter.
    """

    @classmethod
    def INPUT_TYPES(cls):
        return {
            "required": {
                "image": ("IMAGE",),
                "preset": ("LR_PRESET",),
                "strength": ("FLOAT", {
                    "default": 1.0,
                    "min": 0.0,
                    "max": 1.0,
                    "step": 0.05,
                    "display": "slider",
                }),
            },
        }

    RETURN_TYPES = ("IMAGE",)
    RETURN_NAMES = ("image",)
    FUNCTION = "apply_preset"
    CATEGORY = "image/filter"

    def apply_preset(self, image, preset, strength):
        if not preset or "settings" not in preset or not preset["settings"]:
            print("[LRTemplate] No valid preset provided, returning original image")
            return (image,)

        settings = preset["settings"]
        preset_name = preset.get("name", "Unknown")

        batch_size = image.shape[0]
        result_images = []

        for i in range(batch_size):
            img_np = image[i].cpu().numpy().astype(np.float32)
            adjusted = apply_preset(img_np, settings, strength)
            result_images.append(adjusted)

        result = np.stack(result_images, axis=0)
        result_tensor = torch.from_numpy(result).float()

        print(f"[LRTemplate] Applied preset '{preset_name}' with strength {strength:.2f}")
        return (result_tensor,)


class LRTemplateManual:
    """
    Manually apply Lightroom-style adjustments without a preset file.
    Provides sliders for all major adjustment parameters.
    """

    @classmethod
    def INPUT_TYPES(cls):
        return {
            "required": {
                "image": ("IMAGE",),
                "exposure": ("FLOAT", {"default": 0.0, "min": -5.0, "max": 5.0, "step": 0.05, "display": "slider"}),
                "contrast": ("INT", {"default": 0, "min": -100, "max": 100, "step": 1, "display": "slider"}),
                "highlights": ("INT", {"default": 0, "min": -100, "max": 100, "step": 1, "display": "slider"}),
                "shadows": ("INT", {"default": 0, "min": -100, "max": 100, "step": 1, "display": "slider"}),
                "whites": ("INT", {"default": 0, "min": -100, "max": 100, "step": 1, "display": "slider"}),
                "blacks": ("INT", {"default": 0, "min": -100, "max": 100, "step": 1, "display": "slider"}),
                "clarity": ("INT", {"default": 0, "min": -100, "max": 100, "step": 1, "display": "slider"}),
                "vibrance": ("INT", {"default": 0, "min": -100, "max": 100, "step": 1, "display": "slider"}),
                "saturation": ("INT", {"default": 0, "min": -100, "max": 100, "step": 1, "display": "slider"}),
                "temperature": ("INT", {"default": 0, "min": -100, "max": 100, "step": 1, "display": "slider"}),
                "tint": ("INT", {"default": 0, "min": -100, "max": 100, "step": 1, "display": "slider"}),
                "strength": ("FLOAT", {
                    "default": 1.0, "min": 0.0, "max": 1.0, "step": 0.05, "display": "slider",
                }),
            },
        }

    RETURN_TYPES = ("IMAGE",)
    RETURN_NAMES = ("image",)
    FUNCTION = "apply_manual"
    CATEGORY = "image/filter"

    def apply_manual(self, image, exposure, contrast, highlights, shadows,
                     whites, blacks, clarity, vibrance, saturation,
                     temperature, tint, strength):
        settings = {
            "Exposure2012": exposure,
            "Contrast2012": contrast,
            "Highlights2012": highlights,
            "Shadows2012": shadows,
            "Whites2012": whites,
            "Blacks2012": blacks,
            "Clarity2012": clarity,
            "Vibrance": vibrance,
            "Saturation": saturation,
            "Temperature": temperature,
            "Tint": tint,
        }

        settings = {k: v for k, v in settings.items() if v != 0}

        if not settings:
            return (image,)

        batch_size = image.shape[0]
        result_images = []

        for i in range(batch_size):
            img_np = image[i].cpu().numpy().astype(np.float32)
            adjusted = apply_preset(img_np, settings, strength)
            result_images.append(adjusted)

        result = np.stack(result_images, axis=0)
        result_tensor = torch.from_numpy(result).float()

        return (result_tensor,)


class LRTemplateSearchLoader:
    """
    Load a preset by searching for a filename in a folder (including subfolders).
    Supports .lrtemplate, .xmp, and .cube files.
    Provide the full filename with extension.
    """

    @classmethod
    def INPUT_TYPES(cls):
        return {
            "required": {
                "filename": ("STRING", {
                    "default": "preset.lrtemplate",
                    "multiline": False,
                    "tooltip": "Full filename with extension (.lrtemplate / .xmp / .cube)",
                }),
                "search_folder": ("STRING", {
                    "default": "",
                    "multiline": False,
                    "tooltip": "Root folder path to search (subfolders included)",
                }),
            },
        }

    RETURN_TYPES = ("LR_PRESET", "CUBE_LUT",)
    RETURN_NAMES = ("preset", "cube_lut",)
    FUNCTION = "search_and_load"
    CATEGORY = "image/filter"

    @classmethod
    def IS_CHANGED(cls, filename, search_folder):
        found = cls._find_file(filename, search_folder)
        if found:
            return os.path.getmtime(found)
        return float("nan")

    @staticmethod
    def _find_file(filename: str, search_folder: str):
        if not search_folder or not os.path.isdir(search_folder):
            return None
        for root, dirs, files in os.walk(search_folder):
            if filename in files:
                return os.path.join(root, filename)
        return None

    def search_and_load(self, filename, search_folder):
        empty_preset = {}
        empty_cube = {}

        if not search_folder or not os.path.isdir(search_folder):
            print(f"[Search Loader] Folder not found: {search_folder}")
            return (empty_preset, empty_cube)

        filepath = self._find_file(filename, search_folder)

        if not filepath:
            print(f"[Search Loader] File '{filename}' not found in '{search_folder}' or its subfolders")
            return (empty_preset, empty_cube)

        ext = os.path.splitext(filepath)[1].lower()

        try:
            if ext in _LR_EXTENSIONS:
                result = _load_lr_preset(filepath)
                print(f"[Search Loader] Loaded LR preset: {result['name']} from {filepath}")
                return (result, empty_cube)
            elif ext == ".cube":
                lut_data = parse_cube(filepath)
                print(f"[Search Loader] Loaded CUBE LUT: {lut_data['title']} ({lut_data['type']}, size={lut_data['size']}) from {filepath}")
                return (empty_preset, lut_data)
            else:
                print(f"[Search Loader] Unsupported file format: {ext}")
                return (empty_preset, empty_cube)
        except Exception as e:
            print(f"[Search Loader] Error parsing '{filepath}': {e}")
            return (empty_preset, empty_cube)


# =============================================================================
# CUBE LUT Nodes
# =============================================================================

class CUBELUTLoader:
    """
    Load a .cube LUT file from the presets folder.
    Supports both 1D and 3D CUBE LUT files.
    """

    @classmethod
    def INPUT_TYPES(cls):
        cube_list = cls._get_cube_list()
        if not cube_list:
            cube_list = ["No .cube files found - add .cube files to presets folder"]

        return {
            "required": {
                "cube_file": (cube_list, {"default": cube_list[0]}),
            },
        }

    RETURN_TYPES = ("CUBE_LUT",)
    RETURN_NAMES = ("cube_lut",)
    FUNCTION = "load_cube"
    CATEGORY = "image/filter"

    @classmethod
    def _get_cube_list(cls):
        cubes = list_cube_files(PRESETS_DIR)
        return [name for name, _ in cubes] if cubes else []

    @classmethod
    def IS_CHANGED(cls, cube_file):
        filepath = os.path.join(PRESETS_DIR, cube_file)
        if os.path.exists(filepath):
            return os.path.getmtime(filepath)
        return float("nan")

    def load_cube(self, cube_file):
        filepath = os.path.join(PRESETS_DIR, cube_file)

        if not os.path.exists(filepath):
            print(f"[CUBE LUT] File not found: {filepath}")
            return ({},)

        try:
            lut_data = parse_cube(filepath)
            print(f"[CUBE LUT] Loaded: {lut_data['title']} ({lut_data['type']}, size={lut_data['size']})")
            return (lut_data,)
        except Exception as e:
            print(f"[CUBE LUT] Error parsing '{cube_file}': {e}")
            return ({},)


class CUBELUTApply:
    """
    Apply a CUBE LUT to an image.
    Supports both 1D and 3D LUTs with strength blending.
    """

    @classmethod
    def INPUT_TYPES(cls):
        return {
            "required": {
                "image": ("IMAGE",),
                "cube_lut": ("CUBE_LUT",),
                "strength": ("FLOAT", {
                    "default": 1.0,
                    "min": 0.0,
                    "max": 1.0,
                    "step": 0.05,
                    "display": "slider",
                }),
            },
        }

    RETURN_TYPES = ("IMAGE",)
    RETURN_NAMES = ("image",)
    FUNCTION = "apply_lut"
    CATEGORY = "image/filter"

    def apply_lut(self, image, cube_lut, strength):
        if not cube_lut or "data" not in cube_lut:
            print("[CUBE LUT] No valid LUT provided, returning original image")
            return (image,)

        batch_size = image.shape[0]
        result_images = []

        for i in range(batch_size):
            img_np = image[i].cpu().numpy().astype(np.float32)
            adjusted = apply_cube_lut(img_np, cube_lut, strength)
            result_images.append(adjusted)

        result = np.stack(result_images, axis=0)
        result_tensor = torch.from_numpy(result).float()

        lut_name = cube_lut.get("title", "Unknown")
        print(f"[CUBE LUT] Applied '{lut_name}' with strength {strength:.2f}")
        return (result_tensor,)


class CUBELUTSearchLoader:
    """
    Load a .cube LUT file by searching for a filename in a folder (including subfolders).
    Provide the full filename with extension.
    """

    @classmethod
    def INPUT_TYPES(cls):
        return {
            "required": {
                "filename": ("STRING", {
                    "default": "lut.cube",
                    "multiline": False,
                    "tooltip": "Full .cube filename with extension, e.g. MyLUT.cube",
                }),
                "search_folder": ("STRING", {
                    "default": "",
                    "multiline": False,
                    "tooltip": "Root folder path to search (subfolders included)",
                }),
            },
        }

    RETURN_TYPES = ("CUBE_LUT",)
    RETURN_NAMES = ("cube_lut",)
    FUNCTION = "search_and_load"
    CATEGORY = "image/filter"

    @classmethod
    def IS_CHANGED(cls, filename, search_folder):
        found = cls._find_file(filename, search_folder)
        if found:
            return os.path.getmtime(found)
        return float("nan")

    @staticmethod
    def _find_file(filename: str, search_folder: str):
        if not search_folder or not os.path.isdir(search_folder):
            return None
        for root, dirs, files in os.walk(search_folder):
            if filename in files:
                return os.path.join(root, filename)
        return None

    def search_and_load(self, filename, search_folder):
        if not search_folder or not os.path.isdir(search_folder):
            print(f"[CUBE Search] Folder not found: {search_folder}")
            return ({},)

        filepath = self._find_file(filename, search_folder)

        if not filepath:
            print(f"[CUBE Search] File '{filename}' not found in '{search_folder}' or its subfolders")
            return ({},)

        try:
            lut_data = parse_cube(filepath)
            print(f"[CUBE Search] Loaded: {lut_data['title']} ({lut_data['type']}, size={lut_data['size']}) from {filepath}")
            return (lut_data,)
        except Exception as e:
            print(f"[CUBE Search] Error parsing '{filepath}': {e}")
            return ({},)


# =============================================================================
# Preset Catalog Generator Node
# =============================================================================

_CATALOG_EXTENSIONS = (".lrtemplate", ".xmp", ".cube")
_CATALOG_PREVIEW_WIDTH = 200
_CATALOG_ROWS_PER_SHEET = 300


class PresetCatalogGenerator:
    """
    Scan a folder for all preset files (.lrtemplate / .xmp / .cube),
    apply each to the input image, and generate an Excel catalog
    with preview thumbnails.
    """

    @classmethod
    def INPUT_TYPES(cls):
        return {
            "required": {
                "image": ("IMAGE",),
                "preset_folder": ("STRING", {
                    "default": "",
                    "multiline": False,
                    "tooltip": "Root folder to search for presets (recursive)",
                }),
                "output_excel": ("STRING", {
                    "default": "D:/preset_catalog.xlsx",
                    "multiline": False,
                    "tooltip": "Output Excel file path",
                }),
                "preview_width": ("INT", {
                    "default": 200,
                    "min": 50,
                    "max": 800,
                    "step": 50,
                }),
                "rows_per_sheet": ("INT", {
                    "default": 300,
                    "min": 10,
                    "max": 1000,
                    "step": 10,
                }),
            },
        }

    RETURN_TYPES = ("STRING", "INT", "INT",)
    RETURN_NAMES = ("output_path", "success_count", "fail_count",)
    FUNCTION = "generate_catalog"
    CATEGORY = "image/filter"
    OUTPUT_NODE = True

    def generate_catalog(self, image, preset_folder, output_excel,
                         preview_width, rows_per_sheet):
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XlImage
        from openpyxl.utils import get_column_letter
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        import comfy.model_management

        if not os.path.isdir(preset_folder):
            print(f"[Catalog] Error: folder not found: {preset_folder}")
            return (output_excel, 0, 0)

        # ... (rest stays the same until the inner loop)
        # Get first image from batch as numpy
        img_np = image[0].cpu().numpy().astype(np.float32)

        # Resize for faster processing
        h, w = img_np.shape[:2]
        max_size = 800
        if max(h, w) > max_size:
            pil_tmp = PILImage.fromarray((np.clip(img_np, 0, 1) * 255).astype(np.uint8))
            ratio = max_size / max(h, w)
            pil_tmp = pil_tmp.resize((int(w * ratio), int(h * ratio)), PILImage.LANCZOS)
            img_np = np.array(pil_tmp).astype(np.float32) / 255.0

        # Find all presets
        preset_files = []
        for root, dirs, files in os.walk(preset_folder):
            for fname in sorted(files):
                if any(fname.lower().endswith(ext) for ext in _CATALOG_EXTENSIONS):
                    preset_files.append(os.path.join(root, fname))

        total = len(preset_files)
        print(f"[Catalog] Found {total} preset files in {preset_folder}")

        if total == 0:
            print("[Catalog] No preset files found.")
            return (output_excel, 0, 0)

        # Group presets by their immediate parent directory name
        from collections import defaultdict
        grouped_presets = defaultdict(list)
        for p in preset_files:
            parent_dir = os.path.basename(os.path.dirname(p))
            if not parent_dir:
                parent_dir = "根目录"
            grouped_presets[parent_dir].append(p)

        # Create workbook
        wb = Workbook()
        wb.remove(wb.active)

        success_count = 0
        fail_count = 0
        total_sheets = 0
        start_time = time.time()

        # Temp directory for preview images (needed for repeated wb.save)
        import tempfile
        
        with tempfile.TemporaryDirectory(prefix="preset_catalog_") as temp_dir:
            # Styling
            header_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
            header_align = Alignment(horizontal="center", vertical="center")
            cell_font = Font(name="Arial", size=10)
            cell_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
            center_align = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin"),
            )

            processed_count = 0

            for folder_name, files_in_folder in grouped_presets.items():
                safe_folder = folder_name[:25].replace(":", "").replace("\\", "").replace("/", "").replace("?", "").replace("*", "").replace("[", "(").replace("]", ")")
                files_in_folder.sort()
                
                for idx, preset_path in enumerate(files_in_folder):
                    # Check for ComfyUI interrupt signal
                    comfy.model_management.throw_exception_if_processing_interrupted()
                    
                    sheet_idx = idx // rows_per_sheet
                    row_in_sheet = idx % rows_per_sheet

                    if row_in_sheet == 0:
                        if len(files_in_folder) > rows_per_sheet:
                            sheet_name = f"{safe_folder}[{sheet_idx}]"
                        else:
                            sheet_name = safe_folder
                        
                        sheet_name = sheet_name[:31]
                        base_sheet_name = sheet_name
                        counter = 1
                        while sheet_name in wb.sheetnames:
                            suffix = f"_{counter}"
                            sheet_name = f"{base_sheet_name[:31-len(suffix)]}{suffix}"
                            counter += 1
                            
                        ws = wb.create_sheet(title=sheet_name)
                        # Header row
                        headers = ["#", "Filename", "Preview", "File Path"]
                        col_widths = [8, 35, 30, 80]
                        for col_idx, (hdr, wid) in enumerate(zip(headers, col_widths), 1):
                            cell = ws.cell(row=1, column=col_idx, value=hdr)
                            cell.font = header_font
                            cell.fill = header_fill
                            cell.alignment = header_align
                            cell.border = thin_border
                            ws.column_dimensions[get_column_letter(col_idx)].width = wid
                        ws.row_dimensions[1].height = 30
                        total_sheets += 1

                    filename = os.path.basename(preset_path)
                    excel_row = row_in_sheet + 2  # +1 header, +1 one-based

                    processed_count += 1
                    elapsed = time.time() - start_time
                    avg = elapsed / processed_count
                    remain = avg * (total - processed_count)
                    print(f"  [{processed_count}/{total}] [{folder_name}] {filename} (eta: {remain:.0f}s)", end="")

                    # Apply preset
                    preview_pil = None
                    try:
                        ext = os.path.splitext(preset_path)[1].lower()
                        if ext == ".lrtemplate":
                            parsed = parse_lrtemplate(preset_path)
                            settings = get_settings(parsed)
                            result_np = apply_preset(img_np, settings, 1.0)
                        elif ext == ".xmp":
                            parsed = parse_xmp(preset_path)
                            settings = parsed.get("settings", {})
                            result_np = apply_preset(img_np, settings, 1.0)
                        elif ext == ".cube":
                            lut_data = parse_cube(preset_path)
                            result_np = apply_cube_lut(img_np, lut_data, 1.0)
                        else:
                            raise ValueError(f"Unsupported: {ext}")

                        # Make thumbnail
                        result_clipped = np.clip(result_np, 0, 1)
                        pil_img = PILImage.fromarray((result_clipped * 255).astype(np.uint8))
                        ratio = preview_width / pil_img.width
                        new_h = int(pil_img.height * ratio)
                        preview_pil = pil_img.resize((preview_width, new_h), PILImage.LANCZOS)
                        success_count += 1
                        print(" OK")
                    except Exception as e:
                        fail_count += 1
                        print(f" FAIL: {e}")

                    # Write row
                    c = ws.cell(row=excel_row, column=1, value=idx + 1)
                    c.font = cell_font; c.alignment = center_align; c.border = thin_border
                    
                    c = ws.cell(row=excel_row, column=2, value=filename)
                    c.font = cell_font; c.alignment = cell_align; c.border = thin_border
                    
                    c = ws.cell(row=excel_row, column=3)
                    c.border = thin_border
                    if preview_pil is not None:
                        tmp_path = os.path.join(temp_dir, f"preview_{processed_count}.png")
                        preview_pil.save(tmp_path, format="PNG")
                        xl_img = XlImage(tmp_path)
                        xl_img.width = preview_width
                        xl_img.height = preview_pil.height
                        ws.add_image(xl_img, f"C{excel_row}")
                        ws.row_dimensions[excel_row].height = max(preview_pil.height * 0.75, 20)
                    else:
                        c.value = "Failed"
                        c.font = Font(name="Arial", size=10, color="FF0000")
                        c.alignment = center_align
                        ws.row_dimensions[excel_row].height = 25
                        
                    c = ws.cell(row=excel_row, column=4, value=preset_path)
                    c.font = cell_font; c.alignment = cell_align; c.border = thin_border

                    # Save every 50 rows
                    if processed_count % 50 == 0:
                        try:
                            wb.save(output_excel)
                            print(f"  [Catalog] Saved progress: {processed_count}/{total}")
                        except Exception as save_err:
                            print(f"  [Catalog] Warning: save failed: {save_err}")

            # Final save
            wb.save(output_excel)
            
        elapsed = time.time() - start_time
        print(f"[Catalog] Done! {success_count} ok, {fail_count} failed, {elapsed:.1f}s -> {output_excel}")

        return (output_excel, success_count, fail_count)


# =============================================================================
# Node Registration
# =============================================================================

NODE_CLASS_MAPPINGS = {
    "LRTemplate Loader": LRTemplateLoader,
    "LRTemplate Apply": LRTemplateApply,
    "LRTemplate Manual": LRTemplateManual,
    "LRTemplate Search Loader": LRTemplateSearchLoader,
    "CUBE LUT Loader": CUBELUTLoader,
    "CUBE LUT Apply": CUBELUTApply,
    "CUBE LUT Search Loader": CUBELUTSearchLoader,
    "Preset Catalog Generator": PresetCatalogGenerator,
}

NODE_DISPLAY_NAME_MAPPINGS = {
    "LRTemplate Loader": "LRTemplate Loader 📷",
    "LRTemplate Apply": "LRTemplate Apply 🎨",
    "LRTemplate Manual": "LRTemplate Manual 🎛️",
    "LRTemplate Search Loader": "LRTemplate Search Loader 🔍",
    "CUBE LUT Loader": "CUBE LUT Loader 🎬",
    "CUBE LUT Apply": "CUBE LUT Apply 🎞️",
    "CUBE LUT Search Loader": "CUBE LUT Search Loader 🔎",
    "Preset Catalog Generator": "Preset Catalog Generator 📊",
}
