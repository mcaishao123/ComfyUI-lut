"""
预设预览 Excel 生成脚本

用法:
    python generate_preset_catalog.py <图片路径> <预设文件夹路径> [输出Excel路径]

功能:
    - 递归搜索文件夹中所有 .lrtemplate / .xmp / .cube 文件
    - 将每个预设应用到指定图片上，生成预览效果图
    - 输出 Excel 文件，包含：文件名、效果图、文件路径
    - 每页最多 300 行数据，自动分页
"""

import os
import sys
import io
import time
import traceback

import numpy as np
from PIL import Image
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Import our parsers (same directory)
script_dir = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, script_dir)

from lr_parser import parse_lrtemplate, get_settings, get_preset_name
from lr_adjust import apply_preset
from xmp_parser import parse_xmp
from cube_parser import parse_cube, apply_cube_lut


# =============================================================================
# Constants
# =============================================================================

ROWS_PER_SHEET = 300
PREVIEW_WIDTH = 200   # Preview image width in pixels
SUPPORTED_EXTENSIONS = (".lrtemplate", ".xmp", ".cube")


# =============================================================================
# Preset Discovery
# =============================================================================

def find_all_presets(folder: str) -> list:
    """Recursively find all preset files in folder."""
    presets = []
    for root, dirs, files in os.walk(folder):
        for fname in sorted(files):
            if any(fname.lower().endswith(ext) for ext in SUPPORTED_EXTENSIONS):
                presets.append(os.path.join(root, fname))
    return presets


# =============================================================================
# Apply Preset to Image
# =============================================================================

def load_image(image_path: str) -> np.ndarray:
    """Load an image as float32 numpy array [0,1]."""
    img = Image.open(image_path).convert("RGB")
    # Resize for faster processing (keep aspect ratio)
    max_size = 800
    if max(img.size) > max_size:
        ratio = max_size / max(img.size)
        new_size = (int(img.size[0] * ratio), int(img.size[1] * ratio))
        img = img.resize(new_size, Image.LANCZOS)
    return np.array(img).astype(np.float32) / 255.0


def apply_preset_to_image(image_np: np.ndarray, preset_path: str) -> np.ndarray:
    """Apply a preset file to an image, return the result."""
    ext = os.path.splitext(preset_path)[1].lower()

    if ext == ".lrtemplate":
        parsed = parse_lrtemplate(preset_path)
        settings = get_settings(parsed)
        if not settings:
            raise ValueError("No settings found in lrtemplate")
        return apply_preset(image_np, settings, strength=1.0)

    elif ext == ".xmp":
        parsed = parse_xmp(preset_path)
        settings = parsed.get("settings", {})
        if not settings:
            raise ValueError("No settings found in xmp")
        return apply_preset(image_np, settings, strength=1.0)

    elif ext == ".cube":
        lut_data = parse_cube(preset_path)
        return apply_cube_lut(image_np, lut_data, strength=1.0)

    else:
        raise ValueError(f"Unsupported format: {ext}")


def numpy_to_pil_thumbnail(image_np: np.ndarray, width: int = PREVIEW_WIDTH) -> Image.Image:
    """Convert numpy array to a PIL thumbnail image."""
    img = np.clip(image_np, 0.0, 1.0)
    img = (img * 255).astype(np.uint8)
    pil_img = Image.fromarray(img)
    # Resize to thumbnail
    ratio = width / pil_img.width
    new_h = int(pil_img.height * ratio)
    pil_img = pil_img.resize((width, new_h), Image.LANCZOS)
    return pil_img


# =============================================================================
# Excel Generation
# =============================================================================

def pil_to_xl_image(pil_img: Image.Image) -> XlImage:
    """Convert PIL image to openpyxl Image object."""
    buf = io.BytesIO()
    pil_img.save(buf, format="PNG")
    buf.seek(0)
    xl_img = XlImage(buf)
    return xl_img


def create_sheet_header(ws):
    """Set up column headers and styling for a worksheet."""
    headers = ["序号", "文件名", "效果预览", "文件路径"]
    col_widths = [8, 35, 30, 80]

    header_font = Font(name="微软雅黑", size=12, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 30


def add_row(ws, row_num: int, index: int, filename: str, preview: Image.Image, filepath: str):
    """Add a data row to the worksheet."""
    excel_row = row_num + 1  # +1 for header

    cell_font = Font(name="微软雅黑", size=10)
    cell_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Column A: Index
    cell_a = ws.cell(row=excel_row, column=1, value=index)
    cell_a.font = cell_font
    cell_a.alignment = center_align
    cell_a.border = thin_border

    # Column B: Filename
    cell_b = ws.cell(row=excel_row, column=2, value=filename)
    cell_b.font = cell_font
    cell_b.alignment = cell_align
    cell_b.border = thin_border

    # Column C: Preview image
    cell_c = ws.cell(row=excel_row, column=3)
    cell_c.border = thin_border

    if preview is not None:
        xl_img = pil_to_xl_image(preview)
        # Scale image to fit cell
        img_width_px = preview.width
        img_height_px = preview.height
        # Set cell size to accommodate image
        xl_img.width = PREVIEW_WIDTH
        xl_img.height = img_height_px
        anchor = f"C{excel_row}"
        ws.add_image(xl_img, anchor)
        # Set row height to fit image (convert px to points: ~0.75)
        ws.row_dimensions[excel_row].height = max(img_height_px * 0.75, 20)
    else:
        cell_c.value = "生成失败"
        cell_c.font = Font(name="微软雅黑", size=10, color="FF0000")
        cell_c.alignment = center_align
        ws.row_dimensions[excel_row].height = 25

    # Column D: File path
    cell_d = ws.cell(row=excel_row, column=4, value=filepath)
    cell_d.font = cell_font
    cell_d.alignment = cell_align
    cell_d.border = thin_border


# =============================================================================
# Main
# =============================================================================

def main():
    if len(sys.argv) < 3:
        print("用法: python generate_preset_catalog.py <图片路径> <预设文件夹路径> [输出Excel路径]")
        print()
        print("示例:")
        print('  python generate_preset_catalog.py "D:\\photo.jpg" "D:\\presets" "D:\\output.xlsx"')
        sys.exit(1)

    image_path = sys.argv[1]
    preset_folder = sys.argv[2]
    output_path = sys.argv[3] if len(sys.argv) > 3 else os.path.join(
        os.path.dirname(image_path), "preset_catalog.xlsx"
    )

    # Validate inputs
    if not os.path.isfile(image_path):
        print(f"错误: 图片文件不存在: {image_path}")
        sys.exit(1)

    if not os.path.isdir(preset_folder):
        print(f"错误: 预设文件夹不存在: {preset_folder}")
        sys.exit(1)

    print(f"📷 图片: {image_path}")
    print(f"📁 预设文件夹: {preset_folder}")
    print(f"📊 输出Excel: {output_path}")
    print()

    # Load image
    print("正在加载图片...")
    image_np = load_image(image_path)
    print(f"  图片尺寸: {image_np.shape[1]}x{image_np.shape[0]}")

    # Find all presets
    print("正在搜索预设文件...")
    preset_files = find_all_presets(preset_folder)
    total = len(preset_files)
    print(f"  找到 {total} 个预设文件")

    if total == 0:
        print("未找到任何预设文件，退出。")
        sys.exit(0)

    # Count by type
    counts = {}
    for f in preset_files:
        ext = os.path.splitext(f)[1].lower()
        counts[ext] = counts.get(ext, 0) + 1
    for ext, count in sorted(counts.items()):
        print(f"    {ext}: {count} 个")
    print()

    # Group presets by their immediate parent directory name
    from collections import defaultdict
    grouped_presets = defaultdict(list)
    for p in preset_files:
        parent_dir = os.path.basename(os.path.dirname(p))
        if not parent_dir:
            parent_dir = "根目录"
        grouped_presets[parent_dir].append(p)

    # Create workbook
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    success_count = 0
    fail_count = 0
    total_sheets = 0
    start_time = time.time()
    processed_count = 0

    import tempfile
    
    # Use a context manager to guarantee cleanup even on crash or manual interrupt
    with tempfile.TemporaryDirectory(prefix="preset_catalog_") as temp_dir:
        for folder_name, files_in_folder in grouped_presets.items():
            # Sanitize folder name for Excel sheet (max 31 chars, no brackets usually but we'll manage length)
            safe_folder = folder_name[:25].replace(":", "").replace("\\", "").replace("/", "").replace("?", "").replace("*", "").replace("[", "(").replace("]", ")")
            
            # Sort files in folder
            files_in_folder.sort()
            
            for idx, preset_path in enumerate(files_in_folder):
                sheet_idx = idx // ROWS_PER_SHEET
                row_in_sheet = idx % ROWS_PER_SHEET
                
                if row_in_sheet == 0:
                    # Create a new sheet
                    if len(files_in_folder) > ROWS_PER_SHEET:
                        sheet_name = f"{safe_folder}[{sheet_idx}]"
                    else:
                        sheet_name = safe_folder
                    
                    # Excel sheet names must be unique and <= 31 chars
                    sheet_name = sheet_name[:31]
                    base_sheet_name = sheet_name
                    counter = 1
                    while sheet_name in wb.sheetnames:
                        suffix = f"_{counter}"
                        sheet_name = f"{base_sheet_name[:31-len(suffix)]}{suffix}"
                        counter += 1
                    
                    ws = wb.create_sheet(title=sheet_name)
                    create_sheet_header(ws)
                    total_sheets += 1

                filename = os.path.basename(preset_path)
                processed_count += 1
                elapsed = time.time() - start_time
                avg_time = elapsed / processed_count
                remaining = avg_time * (total - processed_count)
                print(f"  [{processed_count}/{total}] [{folder_name}] {filename}  (预计剩余: {remaining:.0f}秒)", end="")

                try:
                    result_np = apply_preset_to_image(image_np, preset_path)
                    preview = numpy_to_pil_thumbnail(result_np)
                    success_count += 1
                    print("  ✅")
                except Exception as e:
                    preview = None
                    fail_count += 1
                    print(f"  ❌ {e}")

                # Add to Excel
                excel_row = row_in_sheet + 2  # +1 header, +1 for 0-index

                cell_font = Font(name="微软雅黑", size=10)
                cell_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
                center_align = Alignment(horizontal="center", vertical="center")
                thin_border = Border(
                    left=Side(style="thin"), right=Side(style="thin"),
                    top=Side(style="thin"), bottom=Side(style="thin"),
                )

                c = ws.cell(row=excel_row, column=1, value=idx + 1)
                c.font = cell_font; c.alignment = center_align; c.border = thin_border
                
                c = ws.cell(row=excel_row, column=2, value=filename)
                c.font = cell_font; c.alignment = cell_align; c.border = thin_border
                
                c = ws.cell(row=excel_row, column=3)
                c.border = thin_border
                if preview is not None:
                    tmp_path = os.path.join(temp_dir, f"preview_{processed_count}.png")
                    preview.save(tmp_path, format="PNG")
                    xl_img = XlImage(tmp_path)
                    xl_img.width = PREVIEW_WIDTH
                    xl_img.height = preview.height
                    ws.add_image(xl_img, f"C{excel_row}")
                    ws.row_dimensions[excel_row].height = max(preview.height * 0.75, 20)
                else:
                    c.value = "生成失败"
                    c.font = Font(name="微软雅黑", size=10, color="FF0000")
                    c.alignment = center_align
                    ws.row_dimensions[excel_row].height = 25
                    
                c = ws.cell(row=excel_row, column=4, value=preset_path)
                c.font = cell_font; c.alignment = cell_align; c.border = thin_border

                # Save progress contextually
                if processed_count % 50 == 0:
                    try:
                        wb.save(output_path)
                        print(f"  💾 已保存进度: {processed_count}/{total}")
                    except Exception as save_err:
                        print(f"  ⚠️ 保存失败: {save_err}")

        # Final save
        print()
        print(f"正在保存 Excel 到: {output_path}")
        wb.save(output_path)

    elapsed = time.time() - start_time
    print()
    print("=" * 60)
    print(f"✅ 完成! 总耗时: {elapsed:.1f} 秒")
    print(f"   总预设数: {total}")
    print(f"   成功: {success_count}")
    print(f"   失败: {fail_count}")
    print(f"   总页数: {total_sheets}")
    print(f"   输出: {output_path}")
    print("=" * 60)


if __name__ == "__main__":
    main()
